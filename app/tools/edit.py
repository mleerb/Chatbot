import io
from pathlib import Path

from docx import Document
from langchain_core.tools import tool
from pydantic import BaseModel, Field

from app.services.sessions import Download, sessions, sessions_lock
from app.services.extract import file_context
from app.core.file_types import TEXT_EXTENSIONS


class FileEdit(BaseModel):
    old_text: str = Field(description="Exact text to find in the uploaded file")
    new_text: str = Field(description="Replacement text")
    file_id: str = Field(default="", description="ID of the uploaded file to edit; required when multiple files are uploaded")


def replace_text(data: bytes, filename: str, old_text: str, new_text: str) -> bytes:
    if not old_text:
        raise ValueError("Provide non-empty text to replace.")
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Uploaded text file is not valid UTF-8.") from exc
        if text.count(old_text) != 1:
            raise ValueError("Text to replace must occur exactly once in the uploaded file.")
        updated = text.replace(old_text, new_text, 1)
        return (b"\xef\xbb\xbf" if data.startswith(b"\xef\xbb\xbf") else b"") + updated.encode("utf-8")
    if suffix != ".docx":
        raise ValueError("Editing is supported only for UTF-8 text and .docx files.")
    try:
        document = Document(io.BytesIO(data))
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        matches = [paragraph for paragraph in paragraphs if old_text in paragraph.text]
        if len(matches) != 1 or matches[0].text.count(old_text) != 1 or not matches[0].runs:
            raise ValueError("Text to replace must occur exactly once within one paragraph or table cell.")
        paragraph = matches[0]
        runs = list(paragraph.runs)
        start = paragraph.text.index(old_text)
        end = start + len(old_text)
        position = 0
        inserted = False
        for run in runs:
            run_start = position
            run_end = position + len(run.text)
            position = run_end
            if run_end <= start or run_start >= end:
                continue
            left = run.text[:max(0, start - run_start)] if run_start <= start else ""
            right = run.text[max(0, end - run_start):] if run_end >= end else ""
            run.text = left + (new_text if not inserted else "") + right
            inserted = True
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not edit this Word file.") from exc


def create_edit_tool(token: str):
    @tool(
        args_schema=FileEdit,
        description="Edit an uploaded UTF-8 text or Word .docx file by ID, replacing one exact, unique passage. Retains every edited revision for download and never changes the original. Use only when the user requests an edit, never for binary files or images.",
    )
    def edit_uploaded_file(old_text: str, new_text: str, file_id: str = "") -> str:
        with sessions_lock:
            session = sessions.get(token)
        if session is None:
            return "Session expired; no edit saved."
        if not session.uploads:
            return "No uploaded file to edit."
        if not file_id and len(session.uploads) > 1:
            return "Choose the uploaded file_id to edit; multiple files are uploaded."
        item = session.uploads.get(file_id) if file_id else session.latest_upload
        if item is None:
            return "Uploaded file not found; no edit saved."
        try:
            edited = replace_text(item.edits[-1].data if item.edits else item.data,
                                  item.filename, old_text, new_text)
        except ValueError as exc:
            return str(exc)
        revision = Download(item.filename, edited)
        item.edits.append(revision)
        item.document, item.image_url, item.status = file_context(item.filename, edited)
        return f"Edited copy ready for download: {item.filename} (file_id={item.id}, revision_id={revision.id}). The original is unchanged."

    return edit_uploaded_file
