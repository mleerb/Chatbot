import io

from docx import Document

from app.formats.blocks import document_blocks, table_rows


def make_docx(content: str) -> bytes:
    document = Document()
    for block in document_blocks(content):
        if not isinstance(block, dict):
            raise ValueError("Document blocks must be objects.")
        kind = block.get("type")
        if kind in {"heading", "paragraph"} and isinstance(block.get("text"), str):
            if kind == "heading":
                document.add_heading(block["text"], level=1)
            else:
                document.add_paragraph(block["text"])
        elif kind == "table":
            rows = table_rows(block)
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for index, row in enumerate(rows):
                for column, cell in enumerate(row):
                    table.cell(index, column).text = str(cell)
        else:
            raise ValueError("Unsupported document block.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
