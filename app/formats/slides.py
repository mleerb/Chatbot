import io

from pptx import Presentation

from app.formats.blocks import structured_content


def make_pptx(content: str) -> bytes:
    slides = structured_content(content, "slides")
    if not slides:
        raise ValueError("Provide at least one slide.")
    presentation = Presentation()
    for slide_data in slides:
        if not isinstance(slide_data, dict) or not isinstance(slide_data.get("title"), str) or not isinstance(slide_data.get("body", ""), str):
            raise ValueError("Each slide needs a title and text body.")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data["title"]
        slide.placeholders[1].text = slide_data.get("body", "")
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
