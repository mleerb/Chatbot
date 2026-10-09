import io

from openpyxl import Workbook
from openpyxl.styles import Font

from app.formats.blocks import structured_content


def make_xlsx(content: str) -> bytes:
    sheets = structured_content(content, "sheets")
    if not sheets:
        raise ValueError("Provide at least one worksheet.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        if not isinstance(sheet, dict) or not isinstance(sheet.get("name"), str) or not isinstance(sheet.get("rows"), list):
            raise ValueError("Each sheet needs a name and a rows array.")
        worksheet = workbook.create_sheet(sheet["name"])
        for index, row in enumerate(sheet["rows"], start=1):
            if not isinstance(row, list) or not all(value is None or isinstance(value, (str, int, float, bool)) for value in row):
                raise ValueError("Sheet rows must contain text, numbers, booleans or null.")
            for column, value in enumerate(row, start=1):
                cell = worksheet.cell(index, column, value)
                if index == 1:
                    cell.font = Font(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
