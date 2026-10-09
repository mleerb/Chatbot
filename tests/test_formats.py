import io
import json
import unittest
import zipfile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.tools.create import validate_file


class FormatTests(unittest.TestCase):
    def test_word_headings_and_tables(self):
        payload = {"blocks": [{"type": "heading", "text": "Report"},
                              {"type": "paragraph", "text": "Summary"},
                              {"type": "table", "rows": [["Name", "Score"], ["A", 3]]}]}
        _, data = validate_file('report.docx', json.dumps(payload))
        document = Document(io.BytesIO(data))
        self.assertEqual(document.paragraphs[0].text, 'Report')
        self.assertEqual(document.paragraphs[0].style.name, 'Heading 1')
        self.assertEqual(document.tables[0].cell(1, 1).text, '3')

    def test_pdf_is_real_pdf(self):
        _, data = validate_file('report.pdf', json.dumps({"blocks": [
            {"type": "heading", "text": "Report"},
            {"type": "paragraph", "text": "Evidence <literal>"},
            {"type": "table", "rows": [["A", "B"], ["one", "two"]]},
        ]}))
        self.assertTrue(data.startswith(b'%PDF-'))
        self.assertIn(b'%%EOF', data[-100:])

    def test_workbook_sheets_and_values(self):
        _, data = validate_file('data.xlsx', json.dumps({"sheets": [
            {"name": "Scores", "rows": [["Name", "Score"], ["Alice", 5]]},
            {"name": "More", "rows": [["Value"], [7]]},
        ]}))
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        self.assertEqual(workbook.sheetnames, ['Scores', 'More'])
        self.assertEqual(workbook['Scores']['B2'].value, 5)
        workbook.close()

    def test_slides(self):
        _, data = validate_file('slides.pptx', json.dumps({"slides": [
            {"title": "Agenda", "body": "Overview"},
            {"title": "Conclusion", "body": "Done"},
        ]}))
        presentation = Presentation(io.BytesIO(data))
        self.assertEqual(len(presentation.slides), 2)
        self.assertEqual(presentation.slides[0].shapes.title.text, 'Agenda')
        self.assertEqual(presentation.slides[1].placeholders[1].text, 'Done')

    def test_zip_and_path_rejection(self):
        _, data = validate_file('bundle.zip', json.dumps({"entries": [
            {"name": "notes/readme.md", "content": "Hi"},
            {"name": "data.csv", "content": "a,b"},
        ]}))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(archive.read('notes/readme.md'), b'Hi')
        for entry in ('../bad.txt', '/bad.txt', 'bad.exe', 'C:bad.txt'):
            with self.subTest(entry=entry), self.assertRaises(ValueError):
                validate_file('bundle.zip', json.dumps({"entries": [{"name": entry, "content": "no"}]}))

    def test_invalid_table_rows_are_rejected_in_word_and_pdf(self):
        for extension in ('docx', 'pdf'):
            for rows in ([], [[]], [['A'], ['B', 'C']], [['A'], [None]]):
                with self.subTest(extension=extension, rows=rows), self.assertRaises(ValueError):
                    validate_file(f'bad.{extension}', json.dumps({
                        'blocks': [{'type': 'table', 'rows': rows}]
                    }))

    def test_invalid_structure_and_unsupported_formats(self):
        for filename, content in [('bad.xlsx', 'not json'), ('bad.pptx', '{}'),
                                  ('bad.docx', '{"blocks":[{"type":"unsupported"}]}'),
                                  ('bad.pdf', '{"blocks":[{"type":"unsupported"}]}'),
                                  ('program.exe', 'MZ'), ('picture.png', 'fake')]:
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                validate_file(filename, content)


if __name__ == '__main__':
    unittest.main()
