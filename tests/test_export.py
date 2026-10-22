import io
import unittest

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.export import markdown_blocks
from app.services.sessions import sessions
from tests.fake_database import mock_database

ANSWER = """# Plan

Some **bold** and `code` with a [link](https://example.com).

| Name | Score |
| --- | ---: |
| Ann | 9 |
| Bo |

- First
  - Nested
1. Step one

```python
print("hi")
```

---
> Quoted
"""


class ExportTests(unittest.TestCase):
    def setUp(self):
        sessions.clear()
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        self.addCleanup(sessions.clear)
        self.database = mock_database(self)

    def export(self, kind, content=ANSWER):
        return self.client.post('/api/export', json={'content': content, 'format': kind})

    def test_markdown_becomes_document_blocks(self):
        self.assertEqual(markdown_blocks(ANSWER), [
            {"type": "heading", "text": "Plan"},
            {"type": "paragraph", "text": "Some bold and code with a link (https://example.com)."},
            {"type": "table", "rows": [["Name", "Score"], ["Ann", "9"], ["Bo", ""]]},
            {"type": "paragraph", "text": "• First"},
            {"type": "paragraph", "text": "    • Nested"},
            {"type": "paragraph", "text": "1. Step one"},
            {"type": "paragraph", "text": 'print("hi")'},
            {"type": "paragraph", "text": "Quoted"},
        ])
        self.assertEqual(markdown_blocks("|---|"), [{"type": "paragraph", "text": ""}])

    def test_markdown_export_keeps_the_answer(self):
        response = self.export('md')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), ANSWER)
        self.assertIn("aster-answer.md", response.headers['content-disposition'])
        self.assertEqual(response.headers['x-content-type-options'], 'nosniff')

    def test_word_export_has_headings_and_tables(self):
        response = self.export('docx')
        self.assertEqual(response.status_code, 200)
        self.assertIn("aster-answer.docx", response.headers['content-disposition'])
        document = Document(io.BytesIO(response.content))
        self.assertEqual(document.paragraphs[0].text, "Plan")
        self.assertEqual(document.tables[0].cell(1, 0).text, "Ann")

    def test_pdf_export_is_a_pdf(self):
        response = self.export('pdf')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_export_validates_input_and_requires_a_session(self):
        self.assertEqual(self.export('html').status_code, 422)
        self.assertEqual(self.export('md', '').status_code, 422)
        self.assertEqual(self.export('md', 'x' * 100_001).status_code, 422)
        self.client.post('/api/auth/logout')
        self.assertEqual(self.export('md').status_code, 401)


if __name__ == '__main__':
    unittest.main()
