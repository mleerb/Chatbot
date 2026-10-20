import io
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.services.sessions import sessions
from app.tools.edit import create_edit_tool, replace_text
from app.main import app
from tests.fake_database import mock_database


class EditTests(unittest.TestCase):
    def setUp(self):
        sessions.clear()
        self.client = TestClient(app)
        mock_database(self)

    def tearDown(self):
        self.client.close()
        sessions.clear()

    def test_text_edit_keeps_original_and_updates_context(self):
        self.client.post('/api/document', files={'file': ('note.md', b'Hello old world\n')})
        token = self.client.cookies.get('chat_session')
        def answer(agent, *args, **kwargs):
            self.assertIn('Edited copy ready', agent.invoke({'old_text': 'old', 'new_text': 'new'}))
            return 'Edited the copy.'
        with patch('app.services.chat.make_agent', side_effect=lambda token, question, sources: create_edit_tool(token)), patch('app.services.chat.answer_question', side_effect=answer):
            response = self.client.post('/api/chat', json={'message': 'Replace old with new'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['uploads'][0]['edits'][0]['filename'], 'note.md')
        self.assertEqual(self.client.get('/api/document/download').content, b'Hello old world\n')
        edited = self.client.get('/api/document/edited')
        self.assertEqual(edited.content, b'Hello new world\n')
        self.assertIn("filename*=UTF-8''note.md", edited.headers['content-disposition'])
        self.assertEqual(edited.headers['x-content-type-options'], 'nosniff')
        self.assertEqual(sessions[token].latest_upload.document, 'Hello new world\n')
        self.assertEqual(self.client.get('/api/state').json()['uploads'][0]['edits'][0]['filename'], 'note.md')
        with TestClient(app) as another:
            another.post('/api/auth/login', json={'email': 'tester@example.com', 'password': 'test-password-123'})
            self.assertEqual(another.get('/api/document/edited').status_code, 404)
        self.client.delete('/api/document')
        self.assertEqual(self.client.get('/api/document/edited').status_code, 404)

    def test_docx_edit_preserves_other_paragraphs_and_runs(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Before ').bold = True
        paragraph.add_run('old ')
        paragraph.add_run('text').italic = True
        document.add_paragraph('Keep this')
        output = io.BytesIO()
        document.save(output)
        original = output.getvalue()
        self.client.post('/api/document', files={'file': ('note.docx', original)})
        token = self.client.cookies.get('chat_session')
        result = create_edit_tool(token).invoke({'old_text': 'old text', 'new_text': 'new text'})
        self.assertIn('Edited copy ready', result)
        edited = Document(io.BytesIO(self.client.get('/api/document/edited').content))
        self.assertEqual(edited.paragraphs[0].text, 'Before new text')
        self.assertEqual(edited.paragraphs[1].text, 'Keep this')
        self.assertTrue(edited.paragraphs[0].runs[0].bold)
        self.assertEqual(self.client.get('/api/document/download').content, original)

    def test_large_text_and_docx_can_be_edited(self):
        text = 'old' + 'x' * 20_001
        self.assertEqual(replace_text(text.encode(), 'note.txt', 'old', 'new').decode(), 'new' + 'x' * 20_001)
        document = Document()
        document.add_paragraph(text)
        output = io.BytesIO()
        document.save(output)
        self.client.post('/api/document', files={'file': ('large.docx', output.getvalue())})
        token = self.client.cookies.get('chat_session')
        self.assertIn('Edited copy ready', create_edit_tool(token).invoke({'old_text': 'old', 'new_text': 'new'}))
        self.assertEqual(sessions[token].latest_upload.document, 'new' + 'x' * 20_001)

    def test_rejects_ambiguous_and_unsupported_edits(self):
        with self.assertRaises(ValueError):
            replace_text(b'old old', 'note.txt', 'old', 'new')
        with self.assertRaises(ValueError):
            replace_text(b'MZ', 'program.exe', 'MZ', 'changed')
        self.client.post('/api/document', files={'file': ('program.exe', b'MZ')})
        token = self.client.cookies.get('chat_session')
        self.assertIn('supported only', create_edit_tool(token).invoke({'old_text': 'MZ', 'new_text': 'Hi'}))
        self.assertEqual(self.client.get('/api/document/edited').status_code, 404)

    def test_upload_preserves_previous_edited_copy(self):
        first = self.client.post('/api/document', files={'file': ('note.txt', b'old')}).json()
        token = self.client.cookies.get('chat_session')
        create_edit_tool(token).invoke({'old_text': 'old', 'new_text': 'new'})
        self.client.post('/api/document', files={'file': ('other.txt', b'other')})
        self.assertEqual(self.client.get('/api/document/edited').status_code, 404)
        self.assertEqual(self.client.get('/api/document/edited', params={'file_id': first['id']}).content, b'new')


if __name__ == '__main__':
    unittest.main()
